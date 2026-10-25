import io

from app.export import to_markdown
from app.export_docs import parse_blocks, to_docx, to_pdf
from app.storage import RecordingMeta

from .conftest import make_transcript

ANALYSIS = {
    "last_template": "riunione",
    "summaries": {"riunione": {"template_name": "Riunione", "title": "T",
                               "markdown": "## Contesto\nRiunione sul **budget** da 40.000 €.\n\n## Punti\n- uno\n  - sotto-punto\n- due"}},
    "keywords": {"keywords": ["budget"], "topics": [{"title": "Budget", "description": "d", "refs": [
        {"timestamp": "00:00:35", "start": 35.0}]}]},
    "extraction": {"decisions": [{"text": "Approvato", "timestamp": "00:00:35"}],
                   "action_items": [{"task": "Preventivo", "owner": "Laura", "due_date": "venerdì", "timestamp": None}],
                   "people": [], "dates": [], "figures": [], "open_questions": []},
}


def _md():
    meta = RecordingMeta(id="x", title="Riunione è ok", created_at="2026-10-25T10:00:00",
                         original_filename="a.m4a", audio_file="audio.m4a", duration=200)
    return to_markdown(meta, make_transcript(6), ANALYSIS)


def test_parse_blocks():
    blocks = list(parse_blocks("# Titolo\n\nRiga uno\nriga due\n\n- a\n  - b\n1. c\n## Sez"))
    assert blocks == [("h", 1, "Titolo"), ("p", "Riga uno riga due"), ("li", 0, "a"), ("li", 1, "b"),
                      ("li", 0, "c"), ("h", 2, "Sez")]


def test_docx_contains_content():
    from docx import Document

    doc = Document(io.BytesIO(to_docx(_md())))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Riunione è ok" in text and "40.000 €" in text and "Preventivo — Laura — entro venerdì" in text
    assert any(r.bold and r.text == "budget" for p in doc.paragraphs for r in p.runs)
    assert doc.paragraphs[0].style.name == "Title"


def test_pdf_is_valid():
    data = to_pdf(_md())
    assert data.startswith(b"%PDF") and len(data) > 2000


def test_pdf_without_truetype_fonts(monkeypatch):
    import app.export_docs as ed

    monkeypatch.setattr(ed, "FONT_CANDIDATES", [])
    assert to_pdf(_md()).startswith(b"%PDF")
