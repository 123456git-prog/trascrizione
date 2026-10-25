"""Export in Word (.docx) e PDF a partire dal Markdown prodotto da export.to_markdown.

Il Markdown è la fonte unica: qui si interpreta il sottoinsieme che usiamo
(titoli, elenchi puntati anche annidati, paragrafi, **grassetto**, _corsivo_).
"""

from __future__ import annotations

import io
import re
from pathlib import Path
from typing import Iterator

from .errors import AppError

INLINE = re.compile(r"(\*\*[^*]+\*\*|_[^_]+_|\*[^*]+\*)")


def parse_blocks(md: str) -> Iterator[tuple]:
    """("h", livello, testo) | ("li", profondità, testo) | ("p", testo)"""
    paragraph: list[str] = []

    def flush():
        if paragraph:
            text = " ".join(line.strip() for line in paragraph)
            paragraph.clear()
            return ("p", text)
        return None

    for raw in md.splitlines():
        line = raw.rstrip()
        heading = re.match(r"^(#{1,6})\s+(.*)", line)
        bullet = re.match(r"^(\s*)[-*+]\s+(.*)", line) or re.match(r"^(\s*)\d+[.)]\s+(.*)", line)
        if heading or bullet or not line.strip():
            block = flush()
            if block:
                yield block
        if heading:
            yield ("h", len(heading.group(1)), heading.group(2).strip())
        elif bullet:
            yield ("li", len(bullet.group(1).replace("\t", "  ")) // 2, bullet.group(2).strip())
        elif line.strip():
            paragraph.append(line)
    block = flush()
    if block:
        yield block


def inline_runs(text: str) -> Iterator[tuple[str, bool, bool]]:
    """(testo, grassetto, corsivo)"""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            yield part[2:-2], True, False
        elif (part.startswith("_") and part.endswith("_")) or (part.startswith("*") and part.endswith("*")):
            yield part[1:-1], False, True
        else:
            yield part, False, False


# --- Word ---------------------------------------------------------------------


def to_docx(md: str) -> bytes:
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
    except ImportError:
        raise AppError("Per l'export Word installa python-docx: pip install -r requirements.txt")

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    first_heading = True
    for block in parse_blocks(md):
        kind = block[0]
        if kind == "h":
            level = block[1]
            if first_heading and level == 1:
                p = doc.add_heading(block[2], level=0)  # titolo del documento
                first_heading = False
            else:
                p = doc.add_heading(block[2], level=min(level - 1 if level > 1 else 1, 4))
            continue
        if kind == "li":
            style = "List Bullet" if block[1] == 0 else f"List Bullet {min(block[1] + 1, 3)}"
            p = doc.add_paragraph(style=style)
            text = block[2]
        else:
            p = doc.add_paragraph()
            text = block[1]
        for chunk, bold, italic in inline_runs(text):
            run = p.add_run(chunk)
            run.bold, run.italic = bold, italic
            if italic and re.fullmatch(r"\(?\d{2}:\d{2}:\d{2}.*\)?", chunk.strip()):
                run.font.color.rgb = RGBColor(0x1F, 0x6F, 0x68)  # timestamp
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


# --- PDF ----------------------------------------------------------------------

FONT_CANDIDATES = [
    # Windows
    ("C:/Windows/Fonts/segoeui.ttf", "C:/Windows/Fonts/segoeuib.ttf", "C:/Windows/Fonts/segoeuii.ttf"),
    ("C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf", "C:/Windows/Fonts/ariali.ttf"),
    # macOS e Linux
    ("/Library/Fonts/Arial.ttf", "/Library/Fonts/Arial Bold.ttf", "/Library/Fonts/Arial Italic.ttf"),
    ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
     "/usr/share/fonts/truetype/dejavu/DejaVuSans-Oblique.ttf"),
    ("/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
     "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf",
     "/usr/share/fonts/truetype/liberation/LiberationSans-Italic.ttf"),
]


def _setup_font(pdf) -> str:
    for regular, bold, italic in FONT_CANDIDATES:
        if Path(regular).exists():
            pdf.add_font("Testo", "", regular)
            pdf.add_font("Testo", "B", bold if Path(bold).exists() else regular)
            pdf.add_font("Testo", "I", italic if Path(italic).exists() else regular)
            return "Testo"
    return ""  # nessun font TrueType trovato: si usano i font base


def _latin(text: str) -> str:
    """Per i font base del PDF (solo caratteri latin-1)."""
    replacements = {"€": "EUR", "’": "'", "‘": "'", "“": '"', "”": '"', "–": "-", "—": "-", "…": "...", "·": "-"}
    for a, b in replacements.items():
        text = text.replace(a, b)
    return text.encode("latin-1", "replace").decode("latin-1")


def to_pdf(md: str) -> bytes:
    try:
        from fpdf import FPDF
    except ImportError:
        raise AppError("Per l'export PDF installa fpdf2: pip install -r requirements.txt")

    pdf = FPDF(format="A4")
    pdf.set_margins(18, 18, 18)
    pdf.set_auto_page_break(True, margin=18)
    family = _setup_font(pdf) or "Helvetica"
    clean = (lambda t: t) if family == "Testo" else _latin
    pdf.add_page()
    width = pdf.w - pdf.l_margin - pdf.r_margin
    sizes = {1: 20, 2: 15, 3: 12.5, 4: 11.5}

    def write_runs(text: str, size: float, indent: float = 0, bullet: str = "") -> None:
        pdf.set_x(pdf.l_margin + indent)
        parts = []
        if bullet:
            parts.append(bullet + " ")
        for chunk, bold, italic in inline_runs(text):
            mark = "**" if bold else ("__" if italic else "")
            parts.append(f"{mark}{clean(chunk).replace('**', '').replace('__', '')}{mark}")
        pdf.set_font(family, "", size)
        pdf.multi_cell(width - indent, size * 0.5, "".join(parts), markdown=True, align="L",
                       new_x="LMARGIN", new_y="NEXT")

    for block in parse_blocks(md):
        kind = block[0]
        if kind == "h":
            level = block[1]
            pdf.ln(4 if level > 1 else 0)
            pdf.set_font(family, "B", sizes.get(level, 11))
            if level == 2:
                pdf.set_text_color(0x1F, 0x6F, 0x68)
            pdf.multi_cell(width, sizes.get(level, 11) * 0.55, clean(block[2]), new_x="LMARGIN", new_y="NEXT")
            pdf.set_text_color(0, 0, 0)
            pdf.ln(1.5)
        elif kind == "li":
            write_runs(block[2], 10.5, indent=4 + 5 * block[1], bullet="•" if family == "Testo" else "-")
            pdf.ln(0.8)
        else:
            write_runs(block[1], 10.5)
            pdf.ln(2)
    return bytes(pdf.output())
